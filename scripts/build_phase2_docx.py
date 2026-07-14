from __future__ import annotations

import csv
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REMOTE_ROOT = Path("/cluster/users/grad/2025/25t8103/project")
ROOT = REMOTE_ROOT if REMOTE_ROOT.exists() else Path(r"I:\Downloads\1\2")
OUT_DIR = ROOT / "phase2_lock_2026-06-11"
if ROOT == REMOTE_ROOT:
    OUT_DIR = ROOT / "paper"
ASSET_DIR = OUT_DIR / "assets"
if ROOT != REMOTE_ROOT and not ASSET_DIR.exists():
    ASSET_DIR = OUT_DIR
EXP = ROOT / "experiments" / "support_v3_2026-06-02"
OUT = Path(os.environ.get("PHASE2_DOCX_OUT", OUT_DIR / "Phase2_Experiment_Report_2026-06-11.docx"))
FIG = ASSET_DIR / "phase2_main_qual_grid_seed12_2026-06-11.png"
FIG2 = ASSET_DIR / "phase2_result_figure2_table1_metrics.png"
FIG3 = ASSET_DIR / "phase2_result_figure3_sd3_tradeoff.png"
FIG4 = ASSET_DIR / "phase2_result_figure4_family_preservation.png"
FIG5 = ASSET_DIR / "phase2_result_figure5_proxy_audit.png"
FIG_EXTERNAL_SD3 = ASSET_DIR / "phase2_external_baseline_bars_sd3.png"
FIG_EXTERNAL_FLUX = ASSET_DIR / "phase2_external_baseline_bars_flux_context.png"
APPENDIX_FIGS = [
    (
        ASSET_DIR / "appendix_grids" / "appendix_T1_attached_accessory_representative_all_methods.png",
        "Figure A1. T1 attached accessory: per-task all-method visual comparison.",
    ),
    (
        ASSET_DIR / "appendix_grids" / "appendix_T2_container_insertion_representative_all_methods.png",
        "Figure A2. T2 container insertion: per-task all-method visual comparison.",
    ),
    (
        ASSET_DIR / "appendix_grids" / "appendix_T3_surface_decal_representative_all_methods.png",
        "Figure A3. T3 surface decal: per-task all-method visual comparison.",
    ),
    (
        ASSET_DIR / "appendix_grids" / "appendix_T4_local_recolor_representative_all_methods.png",
        "Figure A4. T4 local recolor: per-task all-method visual comparison.",
    ),
    (
        ASSET_DIR / "appendix_grids" / "appendix_T5_same_color_material_representative_all_methods.png",
        "Figure A5. T5 same-color material: per-task all-method visual comparison.",
    ),
]
TABLE1_CSV = EXP / "table1_phase2_t1_t5_main_final.csv"
TABLE2A_CSV = EXP / "table2a_phase2_sd3_common_subset_final.csv"
TABLE2B_CSV = EXP / "table2b_phase2_native_context_final.csv"
REGISTRY_CSV = EXP / "phase2_final_selected_runs_2026-06-11.csv"
if ROOT != REMOTE_ROOT:
    if not TABLE1_CSV.exists():
        TABLE1_CSV = OUT_DIR / "table1_phase2_t1_t5_main_final.csv"
    if not TABLE2A_CSV.exists():
        TABLE2A_CSV = OUT_DIR / "table2a_phase2_sd3_common_subset_final.csv"
    if not TABLE2B_CSV.exists():
        TABLE2B_CSV = OUT_DIR / "table2b_phase2_native_context_final.csv"
    if not REGISTRY_CSV.exists():
        REGISTRY_CSV = OUT_DIR / "phase2_final_selected_runs_2026-06-11.csv"


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "EAF2FA"
CALLOUT_FILL = "F3F7FB"


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def metric_table_rows(path: Path) -> list[list[str]]:
    fields = ["label", "n", "outside_l1", "bg_psnr", "bg_lpips_x100", "bg_ssim_luma", "dino_source", "clip_t", "local_clip_t", "edit_score"]
    return [[row.get(field, "") for field in fields] for row in read_csv(path)]


def registry_summary_lines() -> list[str]:
    rows = read_csv(REGISTRY_CSV)
    total = len(rows)
    aligned = sum(1 for row in rows if row.get("metric_status") == "aligned")
    needs = sum(1 for row in rows if row.get("metric_status") == "needs_metric_refresh")
    other = total - aligned - needs
    if not rows:
        return ["Final qualitative figures are resolved through the selected-run registry.", "Registry summary is unavailable in this local build."]
    status = f"The registry contains {total} rows: {aligned} aligned with the current metric CSVs"
    if needs:
        status += f", {needs} flagged as needs_metric_refresh"
    if other:
        status += f", {other} with other non-aligned status"
    status += "."
    return [
        "Final qualitative figures are resolved through phase2_final_selected_runs_2026-06-11.csv, which records the selected run, reason, and metric-alignment status for every family/task/seed/method row.",
        status,
    ]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_grid(table, widths: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc: Document, text: str = "", style: str | None = None, bold_lead: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading("", level=level)
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        size, color = 16, BLUE
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        size, color = 13, BLUE
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        size, color = 12, DARK_BLUE
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_callout(doc: Document, title: str, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_grid(table, [9000])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_run_font(r, size=10)
    doc.add_paragraph()


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.35) -> None:
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        rr = cap.add_run(caption)
        set_run_font(rr, size=9)
    else:
        add_para(doc, f"Figure missing: {path}")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], font_size: int = 9) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_grid(table, widths)
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, GRAY_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            if idx > 0 and len(value) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, size=font_size)
    set_table_grid(table, widths)
    doc.add_paragraph()


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)


def build_doc() -> None:
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("Phase2 Experiment Report")
    set_run_font(r, size=24, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    r = subtitle.add_run("CleanEdit: Non-Edit-Region Preservation First")
    set_run_font(r, size=14, bold=False, color=BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("2026-06-11 | Phase2 T1-T5 | 15 cases x 3 seeds")
    set_run_font(r, size=10)

    add_callout(
        doc,
        "Key Takeaway",
        [
            "The current Phase2 evidence supports a preservation-first claim: CleanEdit best protects the non-edit region while retaining measurable local edit signal.",
            "The result should not be framed as broad general-purpose image-editing superiority or maximum edit-amplitude performance.",
        ],
    )

    add_heading(doc, "1. Experiment Design", 1)
    add_para(doc, "本阶段实验验证的核心不是最大编辑幅度，而是局部编辑时的非编辑区保护。")
    add_para(
        doc,
        "Working claim: CleanEdit prioritizes non-edit-region preservation in localized RF editing while retaining measurable target-local edit signal.",
    )

    add_table(
        doc,
        ["Family", "Current Phase2 Cases"],
        [
            ["T1 attached accessory", "cat_crown; dog_bow_tie_phase2; dog_front_sunglasses_phase2"],
            ["T2 container insertion", "bowl_apple_inside; white_bowl_orange_tabletop_phase2; brown_bowl_lemon_phase2"],
            ["T3 surface decal", "tshirt_star; mug_heart; tote_leaf"],
            ["T4 local recolor", "red_office_chair_to_blue_office_chair; green_mug_orange_phase2; yellow_vase_blue_phase2"],
            ["T5 same-color material", "pillow_same_color_cable_knit; pillow_same_color_cable_knit_grey; pillow_same_color_cable_knit_armchair"],
        ],
        [2600, 6400],
    )

    add_heading(doc, "2. Compared Methods", 1)
    add_table(
        doc,
        ["Method", "Role in Experiment Design"],
        [
            ["RF reconstruction", "Preservation floor / no-target-edit reference"],
            ["Direct target", "Aggressive target guidance baseline"],
            ["Generic support", "Conservative support-control baseline"],
            ["CleanEdit-SD3", "Proposed preservation-first localized RF control"],
        ],
        [2600, 6400],
    )

    add_heading(doc, "3. Metrics", 1)
    add_table(
        doc,
        ["Metric", "Direction", "Interpretation"],
        [
            ["Non-edit MAE", "lower", "Non-edit-region mean absolute error over the fixed non-edit evaluation region; the core preservation metric."],
            ["BG-PSNR", "higher", "PSNR on BG = 1 - dilate(edit mask), aligned with image-editing literature."],
            ["BG-LPIPS x100", "lower", "Perceptual non-edit-region drift after replacing the edit region with source pixels."],
            ["BG-SSIM-luma", "higher", "Masked non-edit-region structure and luminance consistency."],
            ["DINO-source", "higher", "Semantic source identity preservation; not the primary background metric."],
            ["CLIP-T", "higher", "Full-image target prompt adherence; not a preservation metric."],
            ["Local CLIP-T", "higher", "Edit-crop target phrase adherence."],
            ["edit_score", "higher as signal", "Target-vs-source CLIP direction signal; not sufficient alone."],
            ["inside_l1", "descriptive", "Edit-region change magnitude; not monotonic quality."],
        ],
        [1900, 1500, 5600],
    )
    add_para(
        doc,
        "Definition: Non-edit-region mean absolute error, denoted as Non-edit MAE, measures the average pixel deviation between the edited output and the source image over the fixed non-edit evaluation region. We use M_ne = 1 - Dilate(M_edit) for BG metrics to avoid penalizing reasonable boundary/contact changes.",
    )
    add_callout(
        doc,
        "Visual registry note",
        registry_summary_lines(),
    )

    add_heading(doc, "4. Main Qualitative Result Figure", 1)
    add_para(doc, "The main grid uses one representative seed-12 case from each Phase2 family and resolves images through the final selected run registry.")
    add_figure(doc, FIG, "Figure 1. Phase2 qualitative grid: source, direct target, generic support, and CleanEdit.")

    add_heading(doc, "5. Table 1: Phase2 T1-T5 Main Effect", 1)
    add_para(doc, "Scope: 15 task cases x seeds 10/11/12 = 45 rows per method.")
    add_para(doc, "BG metrics use BG = 1 - dilate(edit mask). BG-LPIPS is reported as x100. Inside L1 is descriptive and kept out of the headline table.")
    add_table(
        doc,
        ["Label", "n", "Non-edit MAE", "BG-PSNR", "BG-LPIPSx100", "BG-SSIM", "DINO-src", "CLIP-T", "Local CLIP-T", "edit_score"],
        metric_table_rows(TABLE1_CSV),
        [1700, 450, 800, 800, 950, 800, 800, 800, 950, 800],
        font_size=6,
    )
    add_figure(doc, FIG2, "Figure 2. Phase2 main-effect metrics from Table 1.")
    add_bullet(doc, "CleanEdit-SD3 has the lowest Non-edit MAE and strongest BG-PSNR/BG-LPIPS/BG-SSIM preservation profile.")
    add_bullet(doc, "The positive edit_score indicates it is not merely a no-op preservation method.")
    add_bullet(doc, "Generic support preserves but under-edits; direct target edits aggressively but drifts.")

    add_heading(doc, "6. Table 2a: Same-Backbone SD3 Context", 1)
    add_table(
        doc,
        ["Label", "n", "Non-edit MAE", "BG-PSNR", "BG-LPIPSx100", "BG-SSIM", "DINO-src", "CLIP-T", "Local CLIP-T", "edit_score"],
        metric_table_rows(TABLE2A_CSV),
        [1700, 450, 800, 800, 950, 800, 800, 800, 950, 800],
        font_size=6,
    )
    add_figure(doc, FIG_EXTERNAL_SD3, "Figure 3. Same-backbone SD3 external baseline bar comparison. FlowAlign is excluded from the paper-facing comparison.")
    add_para(
        doc,
        "Interpretation: Some baselines show higher edit scores, but at the cost of weaker non-edit-region preservation. CleanEdit should be presented as preservation-first balance, not maximum edit strength.",
    )

    add_heading(doc, "7. Table 2b: Native RF / FLUX Context", 1)
    add_table(
        doc,
        ["Label", "n", "Non-edit MAE", "BG-PSNR", "BG-LPIPSx100", "BG-SSIM", "DINO-src", "CLIP-T", "Local CLIP-T", "edit_score"],
        metric_table_rows(TABLE2B_CSV),
        [1700, 450, 800, 800, 950, 800, 800, 800, 950, 800],
        font_size=6,
    )
    add_para(doc, "These rows provide native RF / FLUX context and should not be phrased as a strict same-backbone ranking.")
    add_figure(doc, FIG_EXTERNAL_FLUX, "Figure 4. Native/context external baseline bar comparison with CleanEdit-SD3 as a preservation reference.")

    add_heading(doc, "8. Family-Level Preservation", 1)
    add_para(doc, "Family-level Non-edit MAE shows whether the preservation-first trend holds across T1-T5 instead of being driven by a single task type.")
    add_figure(doc, FIG4, "Figure 5. Family-level non-edit-region preservation by method.")

    add_heading(doc, "9. Proxy Internal Audit Precheck", 1)
    add_para(doc, "The current blind-audit sheets are filled with Codex metric-guided proxy ratings. They are not human ratings and must not be described as a user study or independent human visual audit.")
    add_table(
        doc,
        ["Method", "n", "edit", "relation", "preservation", "locality", "artifact", "overall"],
        [
            ["Generic support", "135", "2.044", "2.044", "4.356", "4.533", "2.067", "3.526"],
            ["RF reconstruction", "135", "1.000", "1.000", "4.133", "4.267", "2.067", "2.844"],
            ["Direct target", "135", "1.756", "1.756", "1.244", "1.711", "4.593", "1.489"],
            ["CleanEdit", "135", "4.022", "4.022", "4.756", "4.667", "1.667", "4.667"],
        ],
        [2000, 600, 850, 900, 1250, 950, 850, 900],
        font_size=8,
    )
    add_figure(doc, FIG5, "Figure 6. Proxy internal-audit precheck summary. Ratings are metric-guided proxy scores, not human study results.")
    add_para(doc, "Artifact severity is lower-is-better. The proxy precheck agrees with the preservation-first trend but must not be cited as human evidence.")

    add_heading(doc, "10. Claim Validation", 1)
    add_callout(
        doc,
        "Supported claim",
        [
            "On a controlled Phase2 T1-T5 diagnostic set, CleanEdit improves non-edit-region preservation and source consistency while retaining measurable target-local edit signal.",
        ],
    )
    add_para(doc, "Not supported by the current evidence:")
    for item in [
        "Broad general-purpose image-editing superiority.",
        "Universal superiority under every metric.",
        "Strongest edit amplitude.",
        "A large standalone gain from adaptive feedback alone.",
        "Human visual-audit evidence from the proxy-filled sheets.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "11. Decisions And Next Steps", 1)
    add_number(doc, "Do not enter Phase3 yet; finish Phase2 paper closure first.")
    add_number(doc, "Keep the selected-run registry and metric CSVs frozen together before manuscript submission.")
    add_number(doc, "Human blind audit is optional unless the paper needs human visual-audit claims; proxy scores remain internal precheck evidence only.")
    add_number(doc, "Do not run Phase2 mask sensitivity immediately; keep it as reviewer-defense work.")
    add_number(doc, "Use the preservation-first claim in final writing and avoid old Core-5 / red_chair_blue wording.")

    add_heading(doc, "12. Active Artifacts", 1)
    add_table(
        doc,
        ["Artifact", "Path"],
        [
            ["Scope lock", "PHASE2_LOCK_2026-06-11.md"],
            ["Final selected run registry", "experiments/support_v3_2026-06-02/phase2_final_selected_runs_2026-06-11.csv"],
            ["Markdown report", "paper/phase2_experiment_report_2026-06-11.md"],
            ["Main figure", "paper/assets/phase2_main_qual_grid_seed12_2026-06-11.png"],
            ["Table audit", "experiments/support_v3_2026-06-02/phase2_tables_audit_2026-06-11.json"],
            ["Proxy audit summary", "experiments/support_v3_2026-06-02/blind_internal_audit_phase2_t1_t5_2026-06-11/blind_internal_audit_summary.md"],
        ],
        [2300, 6700],
        font_size=8,
    )

    appendix_section = doc.add_section(WD_SECTION.NEW_PAGE)
    appendix_section.orientation = WD_ORIENT.LANDSCAPE
    appendix_section.page_width = Inches(11)
    appendix_section.page_height = Inches(8.5)
    appendix_section.left_margin = Inches(0.45)
    appendix_section.right_margin = Inches(0.45)
    appendix_section.top_margin = Inches(0.55)
    appendix_section.bottom_margin = Inches(0.55)

    add_heading(doc, "Appendix A. Per-Task All-Method Visual Comparisons", 1)
    add_para(
        doc,
        "Each appendix grid shows the source image plus the formal main and retained external baseline methods: RF reconstruction, Direct, Generic, FlowEdit, SplitFlow, Sam-Flow-SD3, FireFlow, RF-Solver-Edit, ReFlex, Sam-Flow-FLUX, Fixed CleanEdit, and CleanEdit. Seed 12 is used for every Phase2 task case, and image selection follows the final selected run registry.",
    )
    for path, caption in APPENDIX_FIGS:
        add_figure(doc, path, caption, width=10.0)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Phase2 Experiment Report | CleanEdit | 2026-06-11")
    set_run_font(r, size=8)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
